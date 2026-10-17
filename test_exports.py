"""Self-test of the two BOQ download generators (no database changes). Updated v2.4."""
import io, unittest
from openpyxl import load_workbook
from boq_export_v171 import excel_bytes, pdf_bytes

DATA = {
    'drawing_name': 'Civora_Complete_Test_Plan.png',
    'rows': [
        {'work': 'External masonry – indicative', 'items': [
            {'material': 'Bricks', 'quantity': 13542, 'unit': 'nos', 'rate': 13},
            {'material': 'Cement', 'quantity': 34.158, 'unit': 'bags (50 kg)', 'rate': 400},
        ]},
        {'work': 'Internal painting – provisional', 'items': [
            {'material': 'Paint', 'quantity': 33.83, 'unit': 'litres', 'rate': 0},
        ]},
    ],
    'breakdown': [
        {'measurement': 'External wall gross face', 'quantity': 126, 'unit': 'm²'},
        {'measurement': 'Opening deduction', 'quantity': 13.86, 'unit': 'm²'},
    ],
    'notes': ['Opening placement is provisional.'],
    'skipped': ['Wall tiles'],
    'cost_extras': {'labour': 1000, 'transport': 500, 'other': 0, 'tax_percent': 18},
}
MATERIAL_TOTAL = 13542*13 + 34.158*400  # 189709.2


class ExportTests(unittest.TestCase):
    def test_excel(self):
        excel = excel_bytes(DATA)
        self.assertEqual(excel[:2], b'PK')
        book = load_workbook(io.BytesIO(excel), read_only=True)
        # v1.9 added the 'Work Cost Summary' sheet; the old test still expected 3 sheets.
        self.assertEqual(book.sheetnames, ['Material BOQ', 'Work Cost Summary', 'Measurement Basis', 'Assumptions'])
        self.assertAlmostEqual(book['Material BOQ']['F9'].value, MATERIAL_TOTAL, places=2)
        summary = {r[0]: r[1] for r in book['Work Cost Summary'].iter_rows(values_only=True) if r and r[0]}
        subtotal = MATERIAL_TOTAL + 1500
        self.assertAlmostEqual(summary['Subtotal'], subtotal, places=2)
        self.assertAlmostEqual(summary['GRAND TOTAL'], round(subtotal*1.18, 2), places=2)

    def test_measurement_basis_shown(self):
        old = load_workbook(io.BytesIO(excel_bytes(DATA)), read_only=True)
        self.assertEqual(old['Material BOQ']['B4'].value, 'External face (saved before v2.4.1)')
        new = load_workbook(io.BytesIO(excel_bytes({**DATA, 'measurement_basis_label': 'Centreline'})), read_only=True)
        self.assertEqual(new['Material BOQ']['B4'].value, 'Centreline')
        self.assertEqual(new['Measurement Basis']['B1'].value, 'Centreline')

    def test_pdf(self):
        pdf = pdf_bytes(DATA)
        self.assertEqual(pdf[:4], b'%PDF')
        self.assertGreater(len(pdf), 500)


if __name__ == '__main__':
    unittest.main()
